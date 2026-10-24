"""Author the single final comparative Word report from completed, audited local measurements."""
from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
import json
import re
import statistics
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor

LAB = Path(__file__).resolve().parents[1]
OUT = LAB / "deliverables"
BASE = json.loads((LAB / "organizer/report-base.json").read_text())
NAMES = {"short": "Короткий промпт", "detailed": "Техническое задание по ГОСТ", "ordinary": "Обычный длинный промпт"}


def known(value, complete=True, decimals=None):
    rendered = f"{value:.{decimals}f}" if decimals is not None else str(value)
    return rendered if complete else "≥ " + rendered


def para(doc, text, style=None):
    p = doc.add_paragraph(text, style)
    if not style:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def heading(doc, text, level=1):
    p = doc.add_heading(text, level)
    leading = text.split(" ", 1)[0]
    anchor = "arpo_section_" + leading if level == 1 and leading.isdigit() else ("arpo_appendix" if text.startswith("Приложение ") else None)
    if anchor:
        identity = str(len(doc.element.xpath(".//w:bookmarkStart")) + 1)
        start = OxmlElement("w:bookmarkStart");start.set(qn("w:id"), identity);start.set(qn("w:name"), anchor)
        end = OxmlElement("w:bookmarkEnd");end.set(qn("w:id"), identity)
        p._p.insert(1, start);p._p.append(end)
    return p


def navigation(doc):
    p = doc.add_paragraph()
    for index, (label, anchor) in enumerate([
        ("Схема ACOM", "arpo_section_2"), ("Проверки", "arpo_section_4"),
        ("Метрики", "arpo_section_5"), ("ГОСТ", "arpo_section_7"),
        ("Производительность", "arpo_section_8"), ("Исходники", "arpo_section_10"),
        ("Промпты", "arpo_appendix")]):
        if index:
            p.add_run("   ·   ")
        link = OxmlElement("w:hyperlink");link.set(qn("w:anchor"), anchor)
        run = OxmlElement("w:r");props = OxmlElement("w:rPr")
        color = OxmlElement("w:color");color.set(qn("w:val"), "365D83")
        underline = OxmlElement("w:u");underline.set(qn("w:val"), "single")
        props.extend([color, underline]);run.append(props)
        text = OxmlElement("w:t");text.text = label;run.append(text);link.append(run)
        p._p.append(link)
    return p


def external_link(paragraph, text, url):
    target = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink");link.set(qn("r:id"), target)
    run = OxmlElement("w:r");props = OxmlElement("w:rPr")
    color = OxmlElement("w:color");color.set(qn("w:val"), "365D83")
    underline = OxmlElement("w:u");underline.set(qn("w:val"), "single")
    props.extend([color, underline]);run.append(props)
    value = OxmlElement("w:t");value.text = text;run.append(value)
    link.append(run);paragraph._p.append(link)


def table(doc, headers, rows, widths=None):
    if widths:
        # Keep every table inside the A4 text area, including seven-column metrics.
        factor = min(1.0, 16.5 / sum(widths))
        widths = [width * factor for width in widths]
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if widths:
        for column, width in zip(t.columns, widths):
            column.width = Cm(width)
    for i, value in enumerate(headers):
        t.rows[0].cells[i].text = str(value)
    for values in rows:
        for cell, value in zip(t.add_row().cells, values):
            cell.text = str(value)
    for index, row in enumerate(t.rows):
        keep_row = OxmlElement("w:cantSplit");row._tr.get_or_add_trPr().append(keep_row)
        for col, cell in enumerate(row.cells):
            if widths:
                cell.width = Cm(widths[col])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ["top", "left", "bottom", "right"]:
                e = OxmlElement("w:" + edge);e.set(qn("w:val"), "single");e.set(qn("w:sz"), "4");e.set(qn("w:color"), "D9D9D9");borders.append(e)
            tcpr.append(borders)
            margins = OxmlElement("w:tcMar")
            for edge in ["top", "bottom", "left", "right"]:
                e = OxmlElement("w:" + edge);e.set(qn("w:w"), "90");e.set(qn("w:type"), "dxa");margins.append(e)
            tcpr.append(margins)
            if index == 0:
                shade = OxmlElement("w:shd");shade.set(qn("w:fill"), "E7EDF1");tcpr.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name = "Times New Roman";r.font.size = Pt(10.5);r.font.bold = index == 0
                    r.font.color.rgb = RGBColor(0, 0, 0)
    repeat = OxmlElement("w:tblHeader");t.rows[0]._tr.get_or_add_trPr().append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def new_document(subtitle):
    doc = Document()
    section = doc.sections[0]
    doc.settings.odd_and_even_pages_header_footer = False
    section.different_first_page_header_footer = False
    # A4 follows the two laboratory report examples supplied by the user.
    section.page_width = Cm(21);section.page_height = Cm(29.7)
    section.left_margin = Cm(2.3);section.right_margin = Cm(2)
    section.top_margin = Cm(2);section.bottom_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman";normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    language = OxmlElement("w:lang");language.set(qn("w:val"), "ru-RU")
    normal.element.get_or_add_rPr().append(language)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(7)
    for name, size in [("Title", 19), ("Subtitle", 14), ("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 12)]:
        s = doc.styles[name];s.font.name = "Times New Roman";s.font.size = Pt(size);s.font.color.rgb = RGBColor(0, 0, 0)
        s.paragraph_format.space_before = Pt(12);s.paragraph_format.space_after = Pt(7)
    for name in ["Heading 1", "Heading 2", "Heading 3"]:
        doc.styles[name].paragraph_format.keep_with_next = True
    # Remove template theme fonts and its decorative Title border explicitly.
    for style in doc.styles:
        for border in list(style.element.xpath(".//w:pBdr")):
            border.getparent().remove(border)
        properties = style.element.get_or_add_rPr()
        fonts = properties.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts");properties.insert(0, fonts)
        for attribute in list(fonts.attrib):
            if attribute.rsplit("}", 1)[-1].endswith("Theme"):
                del fonts.attrib[attribute]
        for slot in ["ascii", "hAnsi", "eastAsia", "cs"]:
            fonts.set(qn("w:" + slot), "Times New Roman")
    for header in [section.header, section.even_page_header, section.first_page_header]:
        for paragraph in header.paragraphs:
            paragraph.text = ""
    footer = section.footer.paragraphs[0];footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple");field.set(qn("w:instr"), "PAGE");footer._p.append(field)
    para(doc, "ЛАБОРАТОРНАЯ РАБОТА 2").alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = para(doc, BASE["title"], "Title");p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = para(doc, subtitle, "Subtitle");p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, "На основе темы 65 «Блочная сортировка / Bucket sort»").alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, "Выполнила: " + BASE["author"])
    para(doc, "Группа: " + BASE["group"])
    study = Path((LAB / "organizer/study-root.txt").read_text().strip())
    manifest = json.loads((study / "experiment.json").read_text())
    started = datetime.fromisoformat(manifest["created_at"]).astimezone(ZoneInfo("Europe/Moscow")).strftime("%d.%m.%Y")
    prepared = datetime.now(ZoneInfo("Europe/Moscow")).strftime("%d.%m.%Y")
    para(doc, "Измерения начаты: " + started + ". Отчёт подготовлен: " + prepared + ".")
    doc.core_properties.author = BASE["author"]
    doc.core_properties.language = "ru-RU"
    doc.core_properties.title = BASE["title"] + " — " + subtitle
    doc.core_properties.subject = "ACOM containment and aggregation; local Qwen prompt study"
    return doc



def complexity_equation(doc):
    """Native editable Word math with a genuine subscript and superscript."""
    def math_run(text):
        node = OxmlElement("m:r")
        value = OxmlElement("m:t")
        value.text = text
        node.append(value)
        return node

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    math = OxmlElement("m:oMath")
    math.append(math_run("T(n, B) = O(n + B + ∑ "))
    scripts = OxmlElement("m:sSubSup")
    base = OxmlElement("m:e");base.append(math_run("n"))
    subscript = OxmlElement("m:sub");subscript.append(math_run("i"))
    superscript = OxmlElement("m:sup");superscript.append(math_run("2"))
    scripts.extend([base, subscript, superscript])
    math.append(scripts)
    math.append(math_run(");   S(n, B) = O(n + B)"))
    p._p.append(math)
    average = doc.add_paragraph()
    average.alignment = WD_ALIGN_PARAGRAPH.CENTER
    average.paragraph_format.keep_with_next = True
    formula = OxmlElement("m:oMath")
    formula.append(math_run("O(n + B + "))
    fraction = OxmlElement("m:f")
    numerator = OxmlElement("m:num")
    power = OxmlElement("m:sSup")
    base = OxmlElement("m:e");base.append(math_run("n"))
    exponent = OxmlElement("m:sup");exponent.append(math_run("2"))
    power.extend([base, exponent]);numerator.append(power)
    denominator = OxmlElement("m:den");denominator.append(math_run("B"))
    fraction.extend([numerator, denominator]);formula.append(fraction)
    formula.append(math_run(")"))
    average._p.append(formula)
    return p


def common_technical(doc):
    heading(doc, "1 Постановка и связь с первой работой")
    for key in ["purpose", "reuse", "scope"]:
        para(doc, BASE[key])
    heading(doc, "2 Объекты и интеграция ACOM")
    for key in ["containment", "aggregation"]:
        para(doc, BASE[key])
    picture = doc.add_picture(str(LAB / "figures/acom-architecture.png"), width=Cm(16.5))
    picture._inline.docPr.set("descr", "Lab2 включает E и агрегирует Lab65. Lab65 включает A или C и D. D включает B, E агрегирует B. Публичные QI AddRef Release Lab65 делегируют внешнему Unknown Lab2.")
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    para(doc, "Рисунок 1 Состав объектов и механизмы повторного использования")
    for key in ["lifetime", "identity"]:
        para(doc, BASE[key])
    table(doc, ["Интерфейс", "Методы после QI AddRef Release", "Область"], [
        ["IEcoLab65", "SortInt, SortLong, SortFloat, SortDouble, SortLongDouble", "Прежняя сортировка"],
        ["IEcoLab2", "GetVersion возвращает 2", "Новый внешний объект"],
        ["IEcoCalculatorX", "Addition → int32_t; Subtraction → int16_t", "Аргументы int16_t"],
        ["IEcoCalculatorY", "Multiplication → int32_t; Division → int16_t", "Аргументы int16_t"]], [3.4, 8.8, 4.3])
    for key in ["sdk", "vendor", "vendor_identity"]:
        para(doc, BASE[key])
    heading(doc, "3 Алгоритм и сложность")
    para(doc, BASE["sorting"])
    para(doc, "В формуле n — число элементов, B — число блоков, а nᵢ — число элементов блока i. Суммирование ведётся по всем блокам. T обозначает время, S — дополнительную память.")
    complexity_equation(doc)
    para(doc, "Вторая формула описывает ожидаемое время при примерно равномерном заполнении. При концентрации в одном блоке возможен квадратичный рост. Фиксированное B не даёт основания обещать линейность для любого растущего входа. Арифметические обёртки и обращения к фиксированному набору интерфейсов имеют постоянную сложность по n. Дополнительная память объектов ACOM также постоянна по длине массива.")
    para(doc, BASE["domain"])


def testing(doc, audit):
    heading(doc, "4 Независимые проверки и численные результаты")
    para(doc, BASE["testing"]);para(doc, BASE["numeric_testing"])
    rows = []
    for item in audit:
        last = item["validations"][-1]
        counts = [n for n, failures in last["completed_suites"] if failures == "0"]
        rows.append([NAMES[item["label"]], item["repeat"], "/".join(counts), "A и C", "Пройдено"])
    table(doc, ["Промпт", "№", "Проверки по профилям", "Профили", "Итог"], rows, [5.4, 1.5, 4.3, 2.2, 3.1])
    archive_qa = json.loads((LAB / "organizer/package-verification.json").read_text())
    assert archive_qa["passed"] and archive_qa["all_18_profiles_verified"]
    para(doc, "Дополнительно проверена переносимая выдача: три ZIP архива распакованы в чистые папки, контрольные суммы каждого Reuse.c сверены с ответами модели, все девять вариантов заново собраны и проверены в A/C. Все 18 сборок из архивов прошли тот же независимый клиент. Это проверка упаковки после генерации; её время не прибавлялось к затратам модели.")
    para(doc, "Число в таблице — количество выполненных утверждений, включая поэлементные сравнения, а не число разных сценариев. Незавершённый общий шаблон заранее отклонён: отсутствуют createEcoLab65, createEcoLab2 и lab_allocator. Сам контракт и доверенные клиенты прошли предварительную сборку.")
    table(doc, ["Вход или операция", "Результат", "Независимый эталон", "Ошибка"], [
        ["Addition(3, 5)", "8", "Целое 3 + 5", "0"],
        ["Subtraction(7, 2)", "5", "Целое 7 − 2", "0"],
        ["Multiplication(3, 5)", "15", "Целое 3 × 5", "0"],
        ["Division(7, 2)", "3", "Целочисленное 7 / 2", "0"],
        ["7 −4 2 7 0 −100", "−100 −4 0 2 7 7", "Ожидаемый порядок целых", "0"],
        ["5 типов по 257 элементов", "Порядок совпадает", "Поэлементный qsort", "0 на тестах"]], [5.2, 4.2, 4.5, 2.6])
    table(doc, ["Свойство ACOM", "Доказательство в работе", "Граница проверки"], [
        ["Единая идентичность и набор интерфейсов", "Матрицы QI и сравнение canonical Unknown для 4 и 5 интерфейсов в A/C", "Конечные сценарии одного потока"],
        ["Публичное делегирование", "Внешний тестовый контроллер и ручной разбор QI AddRef Release", "Без межпроцессного маршалинга"],
        ["Прямой интерфейс агрегата", "Адрес Lab65 из Lab2 совпадает с результатом private inner QI", "Статическая сборка Eco"],
        ["Включение калькуляторов", "Подмена арифметических слотов поставщика на отличимые результаты", "Выбранные четыре метода"],
        ["Владение и время жизни", "Счётчики ссылок и выделений, удержание X, отказы прямых зависимостей", "Удержание Y отдельно не запускалось"],
        ["Один включённый E", "Ручной разбор: получение X через QI уже приобретённого Y", "Автоматический счётчик проверяет только наличие E"]], [4.6, 7.8, 4.1])
    para(doc, "Первые пять строк — вывод демонстрационного EcoMain, сопоставленный с точными целочисленными значениями и ожидаемым порядком. Независимый клиент дополнительно проверяет 132 пары операндов и массивы по 257 элементов с qsort. Для целых результатов используется точное сравнение, десятичное приближение не нужно. Вещественные регрессии используют представимые целые значения. Нулевая ошибка этих данных не является утверждением о полной проверке всех вещественных входов или расширенной точности long double.")


def experiment(doc, runs, audit, comparison):
    report_reviews = json.loads((LAB / "organizer/report-review-summaries.json").read_text())
    heading(doc, "5 Организация генерации и учёт затрат")
    para(doc, BASE["method"])
    for block in BASE.get("accounting_paragraphs", [BASE["accounting"]]):
        para(doc, block)
    current_study = Path((LAB / "organizer/study-root.txt").read_text().strip())
    calendar_rows = []
    for title, folder in [("Недействительная предварительная серия", LAB / "runs/experiment-20261020-acom-gost"),
                          ("Основная серия и обычный длинный контроль", current_study)]:
        sessions = []
        for run_file in folder.glob("*/run.json"):
            sessions.extend(json.loads(run_file.read_text()).get("sessions", []))
        completed = [s for s in sessions if s.get("started_at") and s.get("ended_at")]
        if completed:
            first = min(datetime.fromisoformat(s["started_at"]) for s in completed)
            last = max(datetime.fromisoformat(s["ended_at"]) for s in completed)
            calendar_rows.append([title,
                                  first.astimezone(ZoneInfo("Europe/Moscow")).strftime("%d.%m.%Y %H:%M:%S"),
                                  last.astimezone(ZoneInfo("Europe/Moscow")).strftime("%d.%m.%Y %H:%M:%S"),
                                  f"{(last - first).total_seconds() / 3600:.2f}"])
    if calendar_rows:
        table(doc, ["Серия", "Начало МСК", "Окончание МСК", "Календарно ч"], calendar_rows,
              [4.5, 4.3, 4.3, 3.4])
        para(doc, "Календарный интервал рассчитан от начала первой до конца последней сохранённой сессии серии. Он включает паузы и чередование попыток, поэтому не равен сумме активного времени и не является отдельным замером труда организатора. Подготовка окружения до первой сессии и оформление отчёта после последней в этот интервал не входят.")
    table(doc, ["Параметр", "Зафиксированное значение"], [
        ["Модель", "mlx-community/Qwen2.5-Coder-14B-Instruct-4bit"],
        ["Ревизия весов", "29efdbab55a161237ab1e432a3abaf6c7ae2b477"],
        ["Ноутбук", "Apple M4 Pro, 24 GiB RAM, macOS arm64"],
        ["ПО", "Python 3.13.7; MLX 0.32.3; mlx-lm 0.32.0"],
        ["ОС и компилятор", "macOS 26.6.2, build 25G83; Apple clang 17.0.0, clang-1700.0.13.3; macOS SDK 15.4"],
        ["Генерация", "temperature 0.7; top_p 0.8; top_k 20; repetition_penalty 1.05; max_tokens 12288"],
        ["Протокол", "single-file, полный Reuse.c; внешняя диагностика"],
        ["Память MLX", "KV 8 bit с начала; prefill 128; кеш промптов 0; лимит wired 12 GiB; защитный порог активной памяти 13 GiB"],
        ["Общий seed SHA256", "172c0e68138870fd19a3fd568985053f7d155357b145b78d1562ca7e18d2697c"]], [4.3, 12.2])
    para(doc, "Дополнительный защитный порог процесса — наблюдаемый peak RSS 16 GiB; активная память MLX проверяется каждые 0,5 с. Это периодическая защита одного процесса, а не жёсткое ограничение всей памяти macOS. Наблюдения памяти и события перезапуска сохраняются отдельно.")
    para(doc, "Обёртка обновлена до 703b25815adfe0c75471a7b9ffb67b3eb980c44c с сохранением локальных улучшений архивирования контекста. Согласование проверено 53 unit тестами и Ruff; eco-doctor и новый базовый каркас прошли 133 проверки ACOM. Для исходников обёртки сохранены контрольные суммы и локальный patch. Установка нового окружения и загрузка весов не понадобились: существующие веса проверены повторно.")
    rows = []
    for r in runs:
        rows.append([NAMES[r["variant"]], r["repeat"], known(r["input_tokens"], r["tokens_complete"]), known(r["output_tokens"], r["tokens_complete"]), known(r["http_requests"], r["requests_complete"]), r["manual_followups"], known(r["active_seconds"], r["timing_complete"], 2)])
    table(doc, ["Промпт", "№", "Входные токены", "Выходные токены", "API", "Уточн.", "Известно активно с"], rows, [4.3, .8, 2.7, 2.7, 1.3, 1.8, 2.9])
    para(doc, BASE["cost"])
    if not all(r["tokens_complete"] for r in runs):
        para(doc, "Учёт токенов и части времени неполон: транспортный сбой краткого варианта не вернул usage; в первом повторе обычного длинного варианта дополнительно прервана незавершённая сессия. Токены и активное время в сводных таблицах — известные суммы. API попытки включают нижнюю границу незарегистрированных прерванных обращений. Исходные события сохранены в requests.jsonl и interrupted-sessions.jsonl. Неполные суммы нельзя считать точными полными затратами.")
    para(doc, BASE["invalid_pilot"])
    pilot = json.loads((LAB / "organizer/invalid-pilot-summary.json").read_text())["totals"]
    warm = json.loads((LAB / "organizer/warmup-accounting.json").read_text())
    total_in = pilot["input_tokens"] + sum(r["input_tokens"] for r in runs) + warm["prompt_tokens"]
    total_out = pilot["output_tokens"] + sum(r["output_tokens"] for r in runs) + warm["completion_tokens"]
    table(doc, ["Объём", "Входные токены", "Выходные токены", "API", "Известно активно с"], [
        ["Предварительная недействительная серия", pilot["input_tokens"], pilot["output_tokens"], pilot["http_requests"], f"{pilot['active_seconds']:.2f}"],
        ["Отдельный прогрев", warm["prompt_tokens"], warm["completion_tokens"], warm["requests"], f"{warm['duration_s']:.3f}"],
        ["Всего для локальной модели обе серии и прогрев", known(total_in, all(r["tokens_complete"] for r in runs)), known(total_out, all(r["tokens_complete"] for r in runs)), known(pilot["http_requests"] + sum(r["http_requests"] for r in runs) + warm["requests"], all(r["requests_complete"] for r in runs)), known(pilot['active_seconds'] + sum(r['active_seconds'] for r in runs) + warm['duration_s'], all(r["timing_complete"] for r in runs), 2)]], [5.7, 3, 3, 1.3, 3.5])
    para(doc, "Расход прогрева указан отдельно. Подготовка окружения, труд организатора и написание отчёта в эти измерения не входят и полностью не измерялись. Предварительные расходы не распределены произвольно между новыми независимыми повторами.")
    heading(doc, "6 История исправлений и анализ полноты")
    para(doc, "Критические ошибки касались самого механизма ACOM. В первом обычном длинном повторе ответ 10 создавал отдельное управление ссылками для публичных X/Y Lab2: арифметические методы были доступны, но независимые профили A/C выявили по 79 нарушений идентичности и переходов. Во втором обычном длинном повторе ответ 7 ошибочно трактовал адрес поля указателя на private inner Unknown как адрес встроенного интерфейса. Такой адрес не является интерфейсом объекта и не может служить canonical Unknown Lab2.")
    para(doc, "Отдельная группа ошибок относится к владению. В ответе 7 второго ГОСТ повтора конструктор агрегата снимал начальную частную ссылку, а обычное создание оставляло лишнюю ссылку при успехе и могло дважды освободить объект при отказе QI. Исправления запрашивались у модели как различение случаев создания, с сохранением всех расходов. Такие дефекты нельзя оценить только по правильности четырёх арифметических результатов.")
    para(doc, "Во втором обычном длинном повторе ответ 8 уже проходил демонстрационный EcoMain, но независимый клиент с UBSan остановился на неверном owner offset в Lab2X_AddRef. Правильные значения калькулятора и даже успешная сборка поэтому проверялись отдельно от корректного управления интерфейсами и ссылками.")
    para(doc, "Соблюдение языка проверялось отдельно от архитектуры. Во втором кратком повторе ответы 7 и 8 сохраняли исправленную логику, но повторно добавляли комментарии C99 вместо допустимых C89 block comments. В первом обычном длинном повторе ответ 13 побайтно повторил предыдущий файл и не устранил четыре отсутствующих предварительных объявления. Эти ответы и расходы учтены как реальные неудачи; формальное наличие подробного уточнения ещё не означает, что модель его выполнила.")
    para(doc, "Ручная проверка ABI дала дополнительный результат: в ответе 7 третьего краткого повтора ECOCALLMETHOD был помещён перед именами в инициализаторах vtables, а не в прототипы и определения callback функций. В используемой ветке SDK для этой платформы макрос пуст, поэтому такой текст компилировался. Требование к объявлениям всё равно не выполнено; исправление запрошено у модели и включено в расходы. Компиляция на одной платформе сама по себе не подтверждает корректное оформление ABI для других платформ.")
    para(doc, "Окончательная оценка понимания основана на коде и наблюдаемом поведении: публичные QI AddRef Release используют контролирующий Unknown; private Unknown отдельно управляет внутренней жизнью; Lab2 отдаёт настоящий интерфейс Lab65, а собственные X/Y передают арифметику одному E. Это подтверждает реализацию заданного механизма после учитываемых уточнений. Прохождение конечных сценариев не означает, что модель самостоятельно верно поняла всю технологию с первого ответа или что измерено её внутреннее понимание.")
    for item in audit:
        heading(doc, NAMES[item["label"]] + " повтор " + str(item["repeat"]), 2)
        v = item["validations"]
        failures = list(dict.fromkeys(x for c in v if not c["passed"] for x in c["failures"]))
        compiler = list(dict.fromkeys(x for c in v if not c["passed"] for x in c["compiler_errors"]))
        para(doc, f"Файл Reuse.c: {item['source_lines']} строк; проверочных циклов: {len(v)}. Первичная независимая проверка: {'пройдена' if v[0]['passed'] else 'не пройдена'}. Окончательная проверка: пройдена. Источник: {item['path']}.")
        if compiler:
            example = compiler[0].split(" [-Werror", 1)[0]
            if len(example) > 180:
                example = example[:177] + "…"
            para(doc, "Пример диагностики неудачной версии: " + example + ". Полные сообщения сохранены в журнале.")
        if failures:
            para(doc, "Отмеченные нарушенные свойства: " + ", ".join(failures[:12]) + ". Полный перечень сохранён в журнале.")
        history_file = LAB / "organizer/review-history.json"
        if history_file.exists():
            history = [r for r in json.loads(history_file.read_text()) if r["run"] == item["path"]]
            if history:
                reviewed_numbers = ", ".join(str(r["request"]) for r in history)
                para(doc, "Ручной разбор сохранён для ответов " + reviewed_numbers + ". Полные замечания по каждой версии находятся в review-history.json и сопоставлены с контрольными суммами исходников.")
                # Human-authored Russian summaries keep the report readable;
                # every detailed finding remains in the immutable-source ledger.
                reviewed = report_reviews[item["path"]]
                assert reviewed["source_sha256"] == history[-1]["source_sha256"]
                for paragraph in reviewed["paragraphs"]:
                    para(doc, paragraph)
        root = Path((LAB / "organizer/study-root.txt").read_text().strip()) / item["path"]
        prompts = sorted(root.glob("followup-*.txt"))
        if prompts:
            prompt_range = prompts[0].name if len(prompts) == 1 else prompts[0].name + " — " + prompts[-1].name
            para(doc, "Сохранено уточнений организатора: " + str(len(prompts)) + " (" + prompt_range + "). Полные тексты доступны в журнале запуска; продолжения уже включают помощь после исходного промпта.")
        if item["context_compactions"]:
            para(doc, "Контекст явно сокращался с архивированием предыдущей переписки. Первоначальные требования и накопленные затраты сохранены. Это влияет на поведение и стоимость последующих ответов.")
    review = json.loads((LAB / "organizer/manual-review.json").read_text())
    para(doc, review["summary"])
    for finding in review.get("findings", []):
        if finding["label"] in {r["variant"] for r in runs}:
            para(doc, NAMES[finding["label"]] + ": " + finding["explanation"])
    para(doc, "Полнота оценивается по одинаковым критериям приёмки: сборка C89, сохранение численного файла, реальные зависимости A–E, identity, lifetime, агрегат Lab65 и делегирование арифметики. Процент от произвольно выбранного общего балла качества не вычисляется. Итоговое прохождение одинаковых проверок не означает одинаковой читаемости или оптимальности кода.")


def gost_section(doc, groups):
    heading(doc, "7 Применение ГОСТ и сравнение с обычным длинным текстом")
    para(doc, BASE["gost"]);para(doc, BASE["ordinary"])
    para(doc, "Сравнение разделено на общую начальную фазу с бюджетом одного ответа модели и полный путь до приёмки. Начальная фаза использует только одинаковую автоматическую диагностику. Позже организатор адаптирует уточнения и использует накопленный опыт, поэтому полные затраты отражают также помощь организатора, язык и содержание продолжений.")
    table(doc, ["Раздел ТЗ", "Применение в промпте"], [
        ["Введение и основания", "Лабораторная 2, повторное использование темы 65, фиксированный контракт"],
        ["Назначение", "Составной компонент ACOM с сортировкой и калькуляторами"],
        ["Требования", "ABI, композиция, identity, lifetime, ошибки, C89 и allocator"],
        ["Документация", "Комментарии о владении; Word пишет организатор"],
        ["Показатели", "Учёт измеренных токенов, HTTP, времени и ограничений"],
        ["Стадии и приёмка", "Порядок разработки и одинаковые независимые проверки"]], [5.2, 11.3])
    para(doc, "До начала измерений точный токенизатор показал вход первого запроса вместе с общим контекстом: 12 252 токена для короткого промпта, 15 515 для ТЗ, 15 293 для обычного длинного. Структура ТЗ добавляет 222 токена к обычному длинному начальному контексту. Эти значения рассчитываются точным токенизатором и сверяются с usage первых запросов; они не заменяют полный расход последующих исправлений.")
    initial_rows = []
    for label in ["short", "detailed", "ordinary"]:
        g = groups[label]["initial_phase"]
        initial_rows.append([NAMES[label], f"{g['passed']}/{g['runs']}", g["input_tokens_median"], g["output_tokens_median"], g["http_requests_median"], f"{g['active_seconds_median']:.2f}"])
    para(doc, "Начальная фаза без уточнений организатора Медианы трёх повторов")
    table(doc, ["Вариант", "Итог", "Вход", "Выход", "API", "Известно активно с"], initial_rows, [4.5, 1.5, 2.7, 2.7, 1.6, 3.5])
    quality = json.loads((LAB / "organizer/initial-quality-gates.json").read_text())
    assert len(quality) == 9 and not any(r["initial_build_passed"] for r in quality)
    para(doc, "Полнота первого ответа проверена отдельно: ни один из девяти ответов не прошёл обязательную сборку C89, то есть 0 из 3 в каждой группе. Диагностика выявила несуществующие имена SDK, попытку подключить справочный Markdown как заголовок, текст пояснений в исходнике и комментарии C99. До исполняемых проверок ACOM эти версии не дошли. Все девять ответов завершены сервером с finish_reason=stop при расходе 2180–3074 выходных токенов и лимите 12 288; принудительного обрезания по этому лимиту не отмечено. Поэтому по первым ответам нельзя утверждать, что один подход уже дал более качественную работающую реализацию. Детали ворот качества сохранены в initial-quality-gates.json.")
    initial_runs = json.loads((LAB / "organizer/measurement-audit/runs.json").read_text())
    paired_rows = []
    for repeat in [1, 2, 3]:
        pair = {r["variant"]: r for r in initial_runs if r["repeat"] == repeat}
        g = pair["detailed"]["initial_phase"]
        o = pair["ordinary"]["initial_phase"]
        paired_rows.append([pair["detailed"]["seed"], g["input_tokens"] - o["input_tokens"], g["output_tokens"] - o["output_tokens"], f"{g['active_seconds'] - o['active_seconds']:+.2f}"])
    para(doc, "Попарные различия первых ответов: ТЗ ГОСТ минус обычный длинный промпт. Отрицательное время означает более быстрый ответ ТЗ в данном повторе. Объём выходного кода зависит от ответа модели; одинаковое содержание требований не даёт одинаковую длину реализации.")
    table(doc, ["Seed", "Разница входных токенов", "Разница выходных токенов", "Разница активного времени с"], paired_rows, [2.1, 4.6, 4.6, 5.2])
    rows = []
    for label in ["short", "detailed", "ordinary"]:
        g = groups[label]
        rows.append([NAMES[label], f"{g['passed']}/{g['runs']}", known(g["input_tokens"]["median"], g["tokens_complete"]), known(g["output_tokens"]["median"], g["tokens_complete"]), known(g["http_requests"]["median"], g["requests_complete"]), g["manual_followups"]["median"], known(g["active_seconds"]["median"], g["timing_complete"], 2)])
    table(doc, ["Вариант", "Итог", "Вход медиана", "Выход медиана", "API мед.", "Уточн. мед.", "Время медиана с"], rows, [4.1, 1.5, 2.5, 2.5, 1.7, 2.1, 2.4])
    token_values = [known(groups[k]["total_tokens"]["median"], groups[k]["tokens_complete"]) for k in ["short", "detailed", "ordinary"]]
    para(doc, "Медианы полного известного расхода входных и выходных токенов: краткий промпт — " + token_values[0] + "; ТЗ по ГОСТ — " + token_values[1] + "; обычный длинный — " + token_values[2] + ". Меньшая медиана времени не означает автоматически меньший расход токенов. Для неполных данных сохранены нижние границы; единый рейтинг по произвольному весу времени и токенов не вводится.")
    picture = doc.add_picture(str(LAB / "figures/generation-time.png"), width=Cm(16.5))
    picture._inline.docPr.set("descr", "Активное время трёх повторов краткого промпта, ТЗ по ГОСТ и обычного длинного текста. Точные значения приведены в таблицах разделов 5 и 7.")
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    para(doc, "Рисунок 2 Известное активное время трёх вариантов с учётом неудач и продолжений")
    faster = min(groups, key=lambda key: groups[key]["active_seconds"]["median"])
    para(doc, "Наименьшая медиана сохранённого активного времени: " + NAMES[faster].lower() + ". Сравнение относится ко всему пути исправлений. Из-за прерванной сессии полные суммарные расходы обычного длинного варианта неизвестны; приведены сохранённые суммы и их нижние границы. Преимущество по одному показателю нельзя автоматически переносить на другие или считать доказанным эффектом только оформления ГОСТ. Начальное содержание длинных вариантов одинаково, а дальнейшие ответы и уточнения адаптивны.")


def performance(doc, labels):
    heading(doc, "8 Измерения готового компонента")
    bench = json.loads((LAB / "organizer/benchmarks-summary.json").read_text())
    para(doc, "Измерения выполнялись после остановки MLX сервера. Доверенный клиент использует clock() и сборку -O2 без UBSan. Арифметика измеряется пакетами по 2 000 000 вызовов четырёх методов; QI(X) и Release — по 500 000 пар. Каждый пакет повторён пять раз. Накопленный результат арифметики проверяется независимо. Ниже медианы по трём независимым реализациям и пяти повторам для каждого профиля.")
    para(doc, "Прямой маршрут — интерфейсы готового E, который сам агрегирует B и поддерживает X и Y. Lab65 использует A/C и D, а Lab2 передаёт арифметику E. Поэтому таблица сравнивает реальные маршруты, а не изолированный одинаковый набор машинных операций. В частности, прямой QI(X) у E уже проходит его собственное агрегирование.")
    routes = {"direct_E": "E напрямую", "Lab65_containment": "Включение Lab65", "Lab2_with_aggregate": "Внешний Lab2"}
    operations = {"arithmetic": "Арифметика", "qi_release": "QI и Release"}
    call_rows = [[NAMES[r["variant"]], r["profile"], routes[r["route"]], operations[r["operation"]], f"{r['median_ns']:.2f}"] for r in bench["calls"] if r["variant"] in labels]
    table(doc, ["Вариант", "A/C", "Маршрут", "Операция", "Медиана нс"], call_rows, [3.8, 1.3, 5.1, 3.8, 2.5])
    para(doc, "Для арифметики показано время одного вызова в смешанном пакете, для QI — время пары QI и Release. Значения описывают выбранный маленький операнд и конкретный компилятор; это не оценка всех затрат приложения. Публичный Lab2 перенаправляет арифметику E, а агрегирование его Lab65 отдельно проверяется матрицей интерфейсов и сортировкой.")
    para(doc, "Сортировка double измерена через реальный агрегированный IEcoLab65 объекта Lab2 и сопоставлена с qsort на одинаковых копиях. Использованы равномерное распределение по допустимому диапазону и узкий обратный кластер, B = 128, по пять повторов; порядок двух сортировок чередовался. Ниже два характерных размера, полная таблица находится в CSV.")
    para(doc, "Входной seed равен 1234567. Для каждого размера массив заново создаётся одинаковым перед каждым повтором. Пять повторов характеризуют вариацию времени на фиксированном наборе данных; это не пять независимых случайных массивов. Создание копий и поэлементная проверка не входят в замер самой сортировки.")
    distributions = {"uniform": "Равномерное", "cluster_reverse": "Обратный кластер"}
    sort_rows = [[NAMES[r["variant"]], r["profile"], distributions[r["distribution"]], r["n"], f"{r['bucket_ms']:.3f}", f"{r['qsort_ms']:.3f}"] for r in bench["sorting"] if r["variant"] in labels and (r["distribution"], r["n"]) in [("uniform", 50000), ("cluster_reverse", 10000)]]
    table(doc, ["Вариант", "A/C", "Распределение", "n", "Bucket мс", "qsort мс"], sort_rows, [3.5, 1.2, 4.1, 1.8, 2.9, 3.0])
    para(doc, "Скорость зависит от распределения элементов по блокам. Узкий обратный кластер вызывает много перемещений вставками, что согласуется с квадратичной оценкой. Совпадение результатов с qsort проверено на каждом измеренном массиве. Стандарт C не задаёт конкретный алгоритм реализации qsort, поэтому сравнение основано на фактических замерах этой платформы.")
    para(doc, "Во всех девяти реализациях файл численного алгоритма один и тот же и сохранён побайтно. Поэтому эти замеры подтверждают интеграцию и характеристики унаследованного Bucket sort. Они не показывают, что один промпт сгенерировал более быстрый алгоритм сортировки: заново генерировалась объектная организация ACOM, а не численные методы.")


def closing(doc, conclusion):
    heading(doc, "9 Выводы и ограничения")
    para(doc, conclusion)
    for block in BASE.get("limits_paragraphs", [BASE["limits"]]):
        para(doc, block)
    para(doc, "Возможные дальнейшие улучшения — уменьшение дублирования однотипных интерфейсных обёрток, более ясные комментарии о владении и отдельное исследование крупных перегруженных блоков сортировки. Они не объявляются выполненными изменениями измеренных вариантов.")
    heading(doc, "10 Файлы и воспроизводимость")
    para(doc, "Исходники предоставлены раздельными архивами короткого, ГОСТ и контрольного обычного длинного вариантов, каждый с тремя повторами. Внутри находятся общий контракт, статические фабрики, неизменные численные методы первой работы, калькуляторы преподавателя, UnitTest, доверенные проверки, промпты и самостоятельный скрипт сборки. Для сборки нужен отдельно полученный Eco SDK и Apple Command Line Tools. Веса модели и окружение солвера в архивы не включены.")
    para(doc, "Полные run.json, requests.jsonl, events.jsonl, followup и архивы истории находятся в labs/acom-reuse-lab2/runs/experiment-20261020-v2-acom-gost. Недействительная предварительная серия отдельно сохранена в experiment-20261020-acom-gost. Измерения сведены в generation-measurements.csv и CSV производительности. Контрольные суммы и результаты итогового аудита приложены отдельно. Репозиторий с результатами второй работы автоматически не публиковался.")
    heading(doc, "Источники")
    for index, (title, source) in enumerate(BASE["sources"], 1):
        p = para(doc, f"{index}. {title}. ")
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if source.startswith(("https://", "http://")):
            external_link(p, source, source)
        else:
            p.add_run(source)


def append_prompts(doc, labels):
    heading(doc, "Приложение Тексты исходных промптов")
    mapping = {"short": "short.txt", "detailed": "detailed.txt", "ordinary": "ordinary-long.txt"}
    for label in labels:
        heading(doc, NAMES[label], 2)
        for line in (LAB / mapping[label]).read_text().strip().splitlines():
            if not line.strip():
                continue
            p = para(doc, line)
            if re.match(r"^\d+(?:\.\d+)*\s", line):
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(4)
                for run in p.runs:
                    run.bold = True


def main():
    audit_root = LAB / "organizer/measurement-audit"
    runs = json.loads((audit_root / "runs.json").read_text())
    audit = json.loads((audit_root / "code-and-validation.json").read_text())
    groups = json.loads((audit_root / "summary.json").read_text())
    assert len(runs) == 9 and all(r["status"] == "passed" for r in runs)
    review = json.loads((LAB / "organizer/manual-review.json").read_text())
    assert review["all_nine_reviewed"] and review["passed"]
    conclusions = json.loads((LAB / "organizer/conclusions.json").read_text())
    OUT.mkdir(exist_ok=True)
    for kind, filename in [("comparison", "report-comparison.docx")]:
        labels = {"short", "detailed", "ordinary"} if kind == "comparison" else {kind}
        subset = [r for r in runs if r["variant"] in labels]
        selected_audit = [r for r in audit if r["label"] in labels]
        doc = new_document("Сравнительный отчёт" if kind == "comparison" else NAMES[kind])
        para(doc, conclusions[kind])
        para(doc, "Основные результаты. Медианы трёх независимых повторов включают неудачные ответы и все продолжения до приёмки.")
        summary_rows = []
        for label in ["short", "detailed", "ordinary"]:
            group = groups[label]
            summary_rows.append([NAMES[label], f"{group['passed']}/{group['runs']}",
                                 known(group["total_tokens"]["median"], group["tokens_complete"]),
                                 known(group["http_requests"]["median"], group["requests_complete"]),
                                 known(group["active_seconds"]["median"] / 60, group["timing_complete"], 2)])
        table(doc, ["Подход", "Итог", "Известные токены медиана", "API медиана", "Известно активно минут медиана"],
              summary_rows, [4.5, 1.5, 3.9, 2.2, 4.4])
        para(doc, "Знак ≥ обозначает нижнюю границу при неполных данных. Один транспортный сбой не вернул usage, у одной прерванной сессии неизвестны usage и длительность; её API попытка учтена отдельно. Сохранённые запросы, время и границы сравнения раскрыты в разделах 5–7. Схема ACOM и интерфейсы — раздел 2, независимые проверки — раздел 4, производительность — раздел 8, файлы и воспроизводимость — раздел 10.")
        navigation(doc)
        doc.add_page_break()
        common_technical(doc);testing(doc, selected_audit);experiment(doc, subset, selected_audit, kind == "comparison")
        gost_section(doc, groups)
        performance(doc, labels)
        closing(doc, conclusions[kind])
        append_prompts(doc, ["short", "detailed"])
        para(doc, "Полный контрольный обычный длинный текст сохранён в ordinary-long.txt и в каждом архиве исходников. Он повторяет те же 19 содержательных абзацев без заголовков ГОСТ; повторная печать одинакового текста в этом отчёте опущена. Проверка эквивалентности и контрольные суммы сохранены в prompt-equivalence-v2.json.")
        doc.save(OUT / filename)
        print(OUT / filename)


if __name__ == "__main__":
    main()
